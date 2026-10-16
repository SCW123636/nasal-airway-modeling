"""生成一个 16:9 但 **0 页** 的有效 PPTX。

为什么：编辑器对 pptx 的页面计数 + upsert-dsl 的「追加」语义组合很挑——
只要编辑器认为文档已有 N 页，追加就会以 N+1 为插入位而越界失败。
唯一稳定可用的是「编辑器认为文档是 0 页」：此时 18 次追加刚好得到 18 页。
"""
import sys

from pptx import Presentation
from pptx.util import Emu

OUT = r"D:\项目\鼻腔建模\outputs\鼻腔功能客观化评估_路演\鼻腔功能客观化评估_路演.pptx"

prs = Presentation()
prs.slide_width = Emu(12192000)   # 16:9，与 DSL 画布 1280x720 同比例
prs.slide_height = Emu(6858000)

xml_slides = prs.slides._sldIdLst
for sld in list(xml_slides):
    rId = sld.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
    prs.part.drop_rel(rId)
    xml_slides.remove(sld)

prs.save(OUT)

chk = Presentation(OUT)
print("saved:", OUT)
print("slide_count:", len(chk.slides._sldIdLst))
print("size(EMU):", chk.slide_width, chk.slide_height)
print("layout count (母版版式仍在):", len(chk.slide_layouts))
