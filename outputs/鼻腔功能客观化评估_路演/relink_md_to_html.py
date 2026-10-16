# -*- coding: utf-8 -*-
"""把 pptx 内超链接目标从 .md 改为 .html（附件已 HTML 化）。

- `附件/….md` → `附件/….html`
- `附件/03-标定与基准计划/`（文件夹，点不开）→ `附件/03-标定与基准计划/索引.html`
- 可重复运行（幂等：已指向 .html 的不再动）。
用法: python relink_md_to_html.py <deck.pptx>
"""
import sys
from pptx import Presentation

deck = sys.argv[1]
prs = Presentation(deck)
n = 0
for slide in prs.slides:
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        for para in shape.text_frame.paragraphs:
            for run in para.runs:
                a = run.hyperlink.address
                if not a:
                    continue
                new = a
                if a == "附件/03-标定与基准计划/":
                    new = "附件/03-标定与基准计划/索引.html"
                elif a.endswith(".md"):
                    new = a[:-3] + ".html"
                if new != a:
                    run.hyperlink.address = new
                    n += 1
prs.save(deck)

# 回读验证
prs2 = Presentation(deck)
links = {}
for i, s in enumerate(prs2.slides, 1):
    for shape in s.shapes:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    a = run.hyperlink.address
                    if a:
                        links.setdefault(i, set()).add(a)
bad = [a for v in links.values() for a in v if a.endswith(".md") or a.endswith("/")]
print(f"{deck}: 改链 {n} 处；残留 .md/文件夹链接 = {len(bad)}")
for p, v in sorted(links.items()):
    print(f"  P{p:02d}: {sorted(v)}")
