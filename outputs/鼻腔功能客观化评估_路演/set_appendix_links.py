# -*- coding: utf-8 -*-
"""给附录页（末页）挂 6 个超链接：按形状文本匹配，对该形状全部 run 挂同一链接 + 下划线。
用法: python set_appendix_links.py <deck.pptx>
"""
import sys
from pptx import Presentation

MAPPING = [
    ("01 · 患者需求调研包", "调研汇总/index.html#cat-01"),
    ("02 · 详细技术链路",   "调研汇总/index.html#cat-02"),
    ("03 · 验证与校准依据", "调研汇总/index.html#cat-03"),
    ("04 · 未来的依据",     "调研汇总/index.html#cat-04"),
    ("05 · 标准训练方法",   "调研汇总/index.html#cat-06"),
    ("06 · 功能演示 Demo",  "http://localhost:8769/workspace/pro/index.html"),
]

def attach(deck):
    prs = Presentation(deck)
    slide = prs.slides[-1]
    hits = []
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        full = shape.text_frame.text
        for key, target in MAPPING:
            if key in full:
                for para in shape.text_frame.paragraphs:
                    for run in para.runs:
                        run.hyperlink.address = target
                        run.font.underline = True
                hits.append(key)
    prs.save(deck)
    # 回读
    prs2 = Presentation(deck)
    links = set()
    for shape in prs2.slides[-1].shapes:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    a = run.hyperlink.address
                    if a:
                        links.add(a)
    print(f"{deck}: 挂载 {len(hits)}/6, 回读 {len(links)} 个唯一链接")
    for l in sorted(links):
        print("   ", l)

if __name__ == "__main__":
    for d in sys.argv[1:]:
        attach(d)
