# -*- coding: utf-8 -*-
"""给 deck 末页（附录页）的四个材料行挂超链接。

用法: python add_hyperlinks.py <deck.pptx>
按「形状级文本匹配」（slidep 会把一行拆成多个 run，run 级匹配会失配），
命中后对该形状内全部 run 挂同一链接并加下划线；整行可点。
"""
import sys
from pptx import Presentation

MAPPING = [
    ("01 · 患者需求调研包", "附件/01-患者需求调研/索引.html"),
    ("02 · 功能演示 Demo",  "http://localhost:8769/workspace/pro/index.html"),
    ("03 · Demo 启动说明",  "附件/02-功能演示Demo/启动说明.html"),
    ("04 · 标定与基准计划", "附件/03-标定与基准计划/从24例到标准_训练设计_v1.html"),
]

deck = sys.argv[1]
prs = Presentation(deck)
slide = prs.slides[-1]
hits = 0
for shape in slide.shapes:
    if not shape.has_text_frame:
        continue
    full = shape.text_frame.text
    for key, target in MAPPING:
        if key in full:
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    run.hyperlink.address = target
                    run.font.underline = True
            hits += 1
prs.save(deck)

# 回读验证
prs2 = Presentation(deck)
links = set()
for shape in prs2.slides[-1].shapes:
    if shape.has_text_frame:
        for para in shape.text_frame.paragraphs:
            for run in para.runs:
                a = run.hyperlink.address
                if a:
                    links.add(a)
print(f"{deck}: 挂载 {hits}/4, 回读 {len(links)} 个唯一链接")
assert links == {t for _, t in MAPPING}, f"回读不符: {links}"
print("VERIFY OK")
