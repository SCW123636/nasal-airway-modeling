# -*- coding: utf-8 -*-
"""
生成 15.3m x 5.3m 现场大屏版（比例约 2.887:1）：
策略 = 居中 16:9 内容 + 两侧延展同色深蓝底（#0A1A2F）+ 克隆原页光晕做侧翼过渡。
- 全幅实心底矩形：拉伸到新画幅（实心无接缝）
- 全幅渐变遮罩：拉伸（纯遮罩，无形变问题）
- 全幅氛围图 PICTURE：同样拉伸铺满新画幅（横向 1.62x）。渲染核验通过：
  ImageGen 氛围图均为抽象发光体（人脸轮廓/网格/光带），横向拉伸视觉成立，
  且铺满比「居中留平翼」在超宽屏上观感更好（2026-10-16 肉眼核对 p01/p05/p13/p14）。
- 其余全部形状：整体右移 delta 居中
- 侧翼各补一枚克隆光晕（左青右琥珀，沿用原页 gradFill alpha；章节扉页无小光晕可克隆则跳过）
"""
import copy
import shutil
from pptx import Presentation
from pptx.oxml.ns import qn

SRC = "剑穹小队 鼻腔_路演.pptx"
DST = "鼻腔功能客观化评估_路演_大屏15.3x5.3.pptx"

H = 6858000                      # 7.5 in = 19.05 cm，保持不变
W_OLD = 12192000                 # 13.333 in
# 新宽 = H * 15.3 / 5.3，取整并保持左右对称
DELTA = (H * 153 - W_OLD * 53) // (2 * 53) + 1   # ~3,802,812
W_NEW = W_OLD + 2 * DELTA                        # 19,797,624 EMU = 54.99 cm

shutil.copyfile(SRC, DST)
prs = Presentation(DST)
prs.slide_width = W_NEW

TOL = 9525  # 1 px


def is_full_canvas(sh):
    try:
        return (sh.left is not None and sh.top is not None
                and sh.left <= TOL and sh.top <= TOL
                and sh.width >= W_OLD - TOL and sh.height >= H - TOL)
    except Exception:
        return False


def find_glow(slide, color):
    """返回包含指定颜色的 gradFill 形状元素（光晕模板）"""
    for sh in slide.shapes:
        xml = sh._element.xml
        if "gradFill" in xml and color in xml and sh.width and sh.width < W_OLD - TOL:
            return sh._element
    return None


_CLONE_ID = [9000]  # 克隆形状专用 ID 段：原样 deepcopy 会复制重复 cNvPr id，WPS 对此拒开文件（2026-10-16）


def place_glow(template_el, cx, cy, size):
    """深拷贝光晕元素，平移到 (cx,cy) 为中心、边长 size；并重编号形状 ID 保证页内唯一"""
    el = copy.deepcopy(template_el)
    cnv = el.find(qn("p:nvSpPr") + "/" + qn("p:cNvPr"))
    if cnv is None:
        cnv = el.find(".//" + qn("p:cNvPr"))
    if cnv is not None:
        _CLONE_ID[0] += 1
        cnv.set("id", str(_CLONE_ID[0]))
    sp_pr = el.find(qn("p:spPr"))
    xfrm = sp_pr.find(qn("a:xfrm"))
    off, ext = xfrm.find(qn("a:off")), xfrm.find(qn("a:ext"))
    off.set("x", str(int(cx - size / 2)))
    off.set("y", str(int(cy - size / 2)))
    ext.set("cx", str(int(size)))
    ext.set("cy", str(int(size)))
    # 旋转归零，避免继承模板角度
    rot = xfrm.get("rot")
    if rot:
        xfrm.set("rot", "0")
    return el


report = []
for idx, slide in enumerate(prs.slides):
    shifted = stretched = 0
    last_full_el = None
    for sh in slide.shapes:
        el = sh._element
        if is_full_canvas(sh):
            sh.left = 0
            sh.width = W_NEW
            stretched += 1
            last_full_el = el
        else:
            sh.left = int(sh.left or 0) + DELTA
            shifted += 1

    # 侧翼光晕：左青右琥珀（找不到对应色就退化为任意光晕）
    added = []
    anchor = last_full_el
    if anchor is not None:
        cyan = find_glow(slide, "22D3EE") or find_glow(slide, "F5A524")
        amber = find_glow(slide, "F5A524") or find_glow(slide, "22D3EE")
        gsize = int(H * 0.72)
        if cyan is not None:
            added.append(place_glow(cyan, DELTA * 0.52, H * 0.36, gsize))
        if amber is not None:
            added.append(place_glow(amber, W_NEW - DELTA * 0.52, H * 0.64, gsize))
        for el in added:
            anchor.addnext(el)
            anchor = el

    report.append(f"P{idx+1:02d} stretch={stretched} shift={shifted} wing_glows={len(added)}")

prs.save(DST)
print(f"W_NEW={W_NEW} EMU = {W_NEW/360000:.2f} cm, H=19.05 cm, ratio={W_NEW/H:.4f} (目标 {15.3/5.3:.4f})")
print(f"DELTA={DELTA}")
print("\n".join(report))
print("SAVED:", DST)
